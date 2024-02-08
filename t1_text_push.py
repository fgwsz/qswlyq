# -*- coding: utf-8 -*-
import os
import docx
from datetime import datetime
import win32clipboard

def get_root_path():
    return os.path.dirname(os.path.abspath(__file__))

def get_current_date():
    date=datetime(datetime.now().year,datetime.now().month,datetime.now().day,0,0,0)
    return date.strftime("%Y%m%d")

def get_file_path():
    return get_root_path()+"\\"+get_current_date()+".docx"

def count_non_empty_cells_in_column(table,column_index):
    count=0
    for row in table.rows:
        cell=row.cells[column_index]
        if cell.text.strip():
            count+=1
    return count

if __name__=="__main__":
    document=docx.Document(get_file_path())
    table=document.tables[1]
    cell_col=1
    cell_row=count_non_empty_cells_in_column(table,cell_col)
    cell=table.cell(cell_row,cell_col)
    # 使用win32clipboard模块获取剪贴板的文本
    win32clipboard.OpenClipboard()
    clipboard_text=win32clipboard.GetClipboardData()
    win32clipboard.CloseClipboard()
    cell.text=clipboard_text
    document.save(get_file_path())
    print(f"{get_file_path()} table[1] index[{cell_row}] 添加文本成功!")
